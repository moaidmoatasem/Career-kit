"""Tolerant reader for LinkedIn's creator-analytics .xlsx export.

The export layout is not documented; this scans every sheet for a header row (first 15 rows)
and maps columns by name, so extra or reordered columns do not matter:
  * a header with "date" plus "impressions"/"engagement"/"followers"  -> daily rows
  * a header with "url" plus "impressions"/"engagement"               -> per-post rows
"""

from __future__ import annotations

import re
import sqlite3
from datetime import date, datetime
from pathlib import Path

from openpyxl import load_workbook

from li_metrics import db

DATE_FORMATS = ("%m/%d/%Y", "%Y-%m-%d", "%d/%m/%Y", "%b %d, %Y")


def _day(value: object) -> str | None:
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, str):
        for fmt in DATE_FORMATS:
            try:
                return datetime.strptime(value.strip(), fmt).date().isoformat()
            except ValueError:
                continue
    return None


def _num(value: object) -> int | None:
    if isinstance(value, bool):
        return None
    if isinstance(value, int | float):
        return int(value)
    if isinstance(value, str) and re.fullmatch(r"[\d,]+", value.strip()):
        return int(value.replace(",", ""))
    return None


def _columns(header: list[object]) -> dict[str, int]:
    cols: dict[str, int] = {}
    for i, cell in enumerate(header):
        name = str(cell or "").strip().lower()
        if "url" in name:
            cols.setdefault("url", i)
        elif "date" in name:
            cols.setdefault("date", i)
        elif "impression" in name:
            cols.setdefault("impressions", i)
        elif "engagement" in name and "rate" not in name:
            cols.setdefault("engagements", i)
        elif "follower" in name:
            cols.setdefault("new_followers", i)
    return cols


def ingest(xlsx: Path, con: sqlite3.Connection) -> dict[str, int]:
    wb = load_workbook(xlsx, read_only=True, data_only=True)
    counts = {"daily": 0, "posts": 0}
    for ws in wb.worksheets:
        rows = list(ws.iter_rows(values_only=True))
        for h, row in enumerate(rows[:15]):
            cols = _columns(list(row))
            metrics = {k: v for k, v in cols.items() if k in db_metrics()}
            if "url" in cols and metrics:
                kind = "posts"
            elif "date" in cols and metrics:
                kind = "daily"
            else:
                continue
            for data in rows[h + 1 :]:
                vals = {k: _num(data[i]) for k, i in metrics.items() if i < len(data)}
                vals = {k: v for k, v in vals.items() if v is not None}
                if kind == "daily":
                    day = _day(data[cols["date"]]) if cols["date"] < len(data) else None
                    if day and vals:
                        db.upsert_daily(con, day, vals)
                        counts["daily"] += 1
                else:
                    url = data[cols["url"]] if cols["url"] < len(data) else None
                    if isinstance(url, str) and url.startswith("http"):
                        pdate = _day(data[cols["date"]]) if "date" in cols else None
                        db.upsert_post(con, url.strip(), pdate, vals)
                        counts["posts"] += 1
            break
    con.commit()
    return counts


def db_metrics() -> set[str]:
    return {"impressions", "engagements", "new_followers"}

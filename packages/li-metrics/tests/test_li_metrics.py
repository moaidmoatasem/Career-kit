from datetime import date, datetime

import pytest
from li_metrics import db, ingest, report
from li_metrics.cli import main
from openpyxl import Workbook


@pytest.fixture
def xlsx(tmp_path):
    wb = Workbook()
    ws = wb.active
    ws.title = "ENGAGEMENT"
    ws.append(["Some title row"])
    ws.append(["Date", "Impressions", "Engagements"])
    for d in range(10, 17):
        ws.append([datetime(2026, 10, d), 100 * d, 10 * d])
    ws.append(["10/03/2026", "1,000", 50])
    ws2 = wb.create_sheet("FOLLOWERS")
    ws2.append(["Date", "New followers"])
    ws2.append([datetime(2026, 10, 16), 5])
    ws3 = wb.create_sheet("TOP POSTS")
    ws3.append(["Post URL", "Post publish date", "Impressions", "Engagement"])
    ws3.append(["https://www.linkedin.com/posts/a", datetime(2026, 10, 14), 900, 40])
    ws3.append(["not a url", None, 1, 1])
    ws4 = wb.create_sheet("JUNK")
    ws4.append(["nothing", "useful"])
    p = tmp_path / "export.xlsx"
    wb.save(p)
    return p


def test_ingest_idempotent_and_report(xlsx, tmp_path):
    con = db.connect(tmp_path / "m.db")
    c1 = ingest.ingest(xlsx, con)
    c2 = ingest.ingest(xlsx, con)
    assert c1 == c2 == {"daily": 9, "posts": 1}
    assert con.execute("SELECT COUNT(*) FROM daily").fetchone()[0] == 8
    assert con.execute("SELECT new_followers FROM daily WHERE day='2026-10-16'").fetchone()[0] == 5
    txt = report.weekly_report(con, date(2026, 10, 16), {"https://www.linkedin.com/posts/a": "QA"})
    assert (
        "Impressions: 9100" in txt and "Top post" in txt and "Best pillar by impressions: QA" in txt
    )
    assert "n/a" in txt  # no prior week


def test_report_empty_and_delta(tmp_path):
    con = db.connect(tmp_path / "e.db")
    assert "engagement rate n/a" in report.weekly_report(con, date(2026, 10, 16))
    db.upsert_daily(con, "2026-10-16", {"impressions": 200, "engagements": 20})
    db.upsert_daily(con, "2026-10-09", {"impressions": 100, "engagements": 10})
    assert "+100%" in report.weekly_report(con, date(2026, 10, 16))
    with pytest.raises(ValueError):
        db.upsert_daily(con, "2026-10-16", {"bad": 1})


def test_helpers():
    assert ingest._day("2026-10-01") == "2026-10-01" and ingest._day("zzz") is None
    assert ingest._day(date(2026, 10, 1)) == "2026-10-01" and ingest._day(5) is None
    assert ingest._num(True) is None and ingest._num("1,234") == 1234 and ingest._num("x") is None
    assert report.load_pillars(None) == {}


def test_cli(xlsx, tmp_path, capsys):
    dbp = tmp_path / "c.db"
    assert main(["ingest", str(xlsx), "--db", str(dbp)]) == 0
    pcsv = tmp_path / "p.csv"
    pcsv.write_text("url,pillar\nhttps://www.linkedin.com/posts/a,QA\n")
    assert main(["report", "--db", str(dbp), "--end", "2026-10-16", "--pillars", str(pcsv)]) == 0
    assert "weekly report" in capsys.readouterr().out
    assert main(["ingest", str(tmp_path / "missing.xlsx"), "--db", str(dbp)]) == 2
